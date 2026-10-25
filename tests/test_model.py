import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

from docx import Document

from vedomost.model import Assembly, Item, aggregate, load_project, save_project
from vedomost.exporter import export_docx


class CalculationTests(unittest.TestCase):
    def test_nested_counts_and_paths(self):
        part=Item(code='D-1',name='Контакт',quantity='3')
        sub=Assembly('S','Сборка',items=[part])
        root=Assembly('R','Изделие',root_quantity='2',items=[Item('Сборочные единицы','S','Сборка','4'),Item(code='D-1',name='Контакт',quantity='1')])
        groups,warnings=aggregate([root,sub])
        g=next(g for g in groups if g.item.code=='D-1')
        self.assertEqual(g.total,26)
        self.assertEqual(g.usages,{'R → S':Decimal(24),'R':Decimal(2)})
        self.assertEqual(warnings,[])

    def test_same_part_in_independent_roots(self):
        groups,_=aggregate([Assembly('A','A',items=[Item(code='D',name='Деталь')]),Assembly('B','B',items=[Item(code='D',name='Деталь',quantity='2')])])
        self.assertEqual(len(groups),1)
        self.assertEqual(groups[0].total,3)
        self.assertEqual(set(groups[0].usages),{'A','B'})

    def test_decimal_materials(self):
        groups,_=aggregate([Assembly('A','A',items=[Item('Материалы','ГОСТ 1','Провод','0,1','м'),Item('Материалы','ГОСТ 1','Провод','0,2','м')])])
        self.assertEqual(groups[0].total,Decimal('0.3'))

    def test_different_dimensions_and_units_are_not_merged(self):
        groups,_=aggregate([Assembly('A','A',items=[Item('Стандартные изделия','ГОСТ 1','Винт М3'),Item('Стандартные изделия','ГОСТ 1','Винт М4'),Item('Материалы','','Провод','1','м'),Item('Материалы','','Провод','1','кг')])])
        self.assertEqual(len(groups),4)

    def test_cycle_without_root(self):
        a=Assembly('A','A',items=[Item('Сборочные единицы','B','B')])
        b=Assembly('B','B',items=[Item('Сборочные единицы','A','A')])
        with self.assertRaisesRegex(ValueError,'Циклическая'):
            aggregate([a,b])

    def test_duplicate_assembly_and_conflicting_parts(self):
        with self.assertRaisesRegex(ValueError,'Две спецификации'):
            aggregate([Assembly('A','A'),Assembly('A','B')])
        with self.assertRaisesRegex(ValueError,'Различаются данные'):
            aggregate([Assembly('A','A',items=[Item(code='D',name='Один'),Item(code='D',name='Другой')])])

    def test_unresolved_child_is_retained(self):
        groups,warnings=aggregate([Assembly('A','A',items=[Item('Сборочные единицы','B','Сборка','2')])])
        self.assertEqual(groups[0].total,2)
        self.assertIn('Не раскрыта сборка B',warnings[0])

    def test_invalid_quantity(self):
        for value in ['', 'NaN', 'Infinity', '-1', '1 м']:
            with self.subTest(value=value), self.assertRaises(ValueError):
                aggregate([Assembly('A','A',items=[Item(name='Деталь',quantity=value)])])

    def test_save_load_and_word(self):
        assemblies=[Assembly('A','Изделие',items=[Item(code='D',name='Деталь')],reviewed=True),Assembly('B','Изделие',items=[Item(code='D',name='Деталь',quantity='2')],reviewed=True)]
        with tempfile.TemporaryDirectory() as tmp:
            project=Path(tmp)/'project.json'
            save_project(project,assemblies)
            self.assertEqual(load_project(project),assemblies)
            groups,_=aggregate(assemblies)
            path=Path(tmp)/'report.docx'
            export_docx(path,groups,'Изделие','ТВК-1')
            doc=Document(path)
            text='\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
            for expected in ['Деталь','Применяемость','A','B','=3']:
                self.assertIn(expected,text)
            footer='\n'.join(c.text for t in doc.sections[0].footer.tables for r in t.rows for c in r.cells)
            self.assertIn('ТВК-1',footer)
            self.assertAlmostEqual(doc.sections[0].page_width.mm,420,places=0)


if __name__=='__main__':
    unittest.main()

"""Editable A3 Word report inspired by the supplied TVK form."""
from pathlib import Path
import textwrap

from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from .model import fmt, SECTIONS

HEADERS = ['№ строки', 'Наименование', 'Обозначение документа\nна поставку', 'Поставщик', 'Применяемость', 'Кол-во', 'Примечание']
WIDTHS = [8, 70, 67, 50, 93, 23, 79]
WRAPS = [3, 32, 30, 23, 43, 10, 36]


def logical_rows(groups):
    rows = []
    section = None
    for group in groups:
        item = group.item
        if item.section != section:
            if rows:
                rows.append(['']*7)
            section = item.section
            rows.append(['', section, '', '', '', '', ''])
        for index, (usage, quantity) in enumerate(group.usages.items()):
            note = ' '.join(x for x in (item.unit if item.unit != 'шт.' else '', item.note) if x)
            cells = ['', item.name if index==0 else '', item.code if index==0 else '', item.supplier if index==0 else '', usage, '' if item.section=='Документация' else fmt(quantity), note]
            parts = [textwrap.wrap(c, width=w, break_long_words=True, break_on_hyphens=False) or [''] for c,w in zip(cells, WRAPS)]
            for n in range(max(map(len,parts))):
                rows.append([p[n] if n<len(p) else '' for p in parts])
        if len(group.usages)>1 and item.section!='Документация':
            rows.append(['','','','','','='+fmt(group.total), item.unit if item.unit!='шт.' else ''])
        rows.append(['']*7)
    return rows


def write_cell(cell, text, size=10, bold=False):
    cell.text = str(text)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(size)
            run.bold = bold


def export_docx(path, groups, title='Изделие', designation='Ведомость ТВК', warnings=()):
    doc = Document()
    normal=doc.styles['Normal']
    normal.font.name='Times New Roman'
    normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(0)
    rows = logical_rows(groups)
    if not rows:
        raise ValueError('Нет позиций для выгрузки.')
    pages=[]
    while rows:
        limit=23 if not pages else 29
        cut=min(limit,len(rows))
        if len(rows)>limit:
            separators=[i+1 for i,row in enumerate(rows[:limit]) if not any(row)]
            if separators and separators[-1]>=limit//2:
                cut=separators[-1]
        pages.append(rows[:cut])
        rows=rows[cut:]
    for page_index, data in enumerate(pages):
        sec=doc.sections[0] if page_index==0 else doc.add_section(WD_SECTION_START.NEW_PAGE)
        sec.orientation=WD_ORIENT.LANDSCAPE
        sec.page_width=Mm(420)
        sec.page_height=Mm(297)
        sec.top_margin=Mm(10)
        sec.bottom_margin=Mm(42 if page_index==0 else 20)
        sec.footer_distance=Mm(8)
        sec.left_margin=Mm(20)
        sec.right_margin=Mm(10)
        borders=OxmlElement('w:pgBorders')
        borders.set(qn('w:offsetFrom'),'page')
        for side in ('top','left','bottom','right'):
            el=OxmlElement('w:'+side)
            for k,v in [('val','single'),('sz','8'),('space','15')]:
                el.set(qn('w:'+k),v)
            borders.append(el)
        sec._sectPr.append(borders)
        limit=23 if page_index==0 else 29
        table=doc.add_table(rows=1, cols=7)
        table.style='Table Grid'
        table.alignment=WD_TABLE_ALIGNMENT.CENTER
        table.autofit=False
        for col,width in zip(table.columns,WIDTHS):
            col.width=Mm(width)
        for cell,width,text in zip(table.rows[0].cells,WIDTHS,HEADERS):
            cell.width=Mm(width)
            write_cell(cell,text,9)
        direction=OxmlElement('w:textDirection')
        direction.set(qn('w:val'),'btLr')
        table.cell(0,0)._tc.get_or_add_tcPr().append(direction)
        table.rows[0].height=Mm(16)
        for i in range(limit):
            row=table.add_row()
            row.height=Mm(7.8)
            row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
            no_split=OxmlElement('w:cantSplit')
            row._tr.get_or_add_trPr().append(no_split)
            values=data[i] if i<len(data) else ['']*7
            values=[str(i+1)]+values[1:]
            for cell,width,value in zip(row.cells,WIDTHS,values):
                cell.width=Mm(width)
                write_cell(cell,value)
                if value in SECTIONS:
                    cell.paragraphs[0].alignment=1
                    for run in cell.paragraphs[0].runs:
                        run.underline=True
        sec.footer.is_linked_to_previous=False
        stamp=sec.footer.add_table(rows=2 if page_index==0 else 1,cols=3,width=Mm(165))
        stamp.style='Table Grid'
        stamp.alignment=WD_TABLE_ALIGNMENT.RIGHT
        stamp.autofit=False
        for col in stamp.columns:
            col.width=Mm(55)
        for cell,text in zip(stamp.rows[0].cells,[designation,f'Лист {page_index+1}',f'Листов {len(pages)}']):
            write_cell(cell,text)
        if page_index==0:
            stamp.rows[1].height=Mm(20)
            write_cell(stamp.cell(1,0),title)
            write_cell(stamp.cell(1,1),'Сводная технологическая\nведомость комплектации')
            write_cell(stamp.cell(1,2),'Разраб. __________\nПров. __________\nУтв. __________')
    if warnings:
        doc.add_page_break()
        doc.add_paragraph('Нераскрытые сборки',style='Heading 1')
        for warning in warnings:
            doc.add_paragraph(warning)
    doc.save(str(Path(path)))
